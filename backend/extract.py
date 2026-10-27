"""Turn a lecture file (.pptx, .ppt or .pdf) into plain text, one entry per slide."""

import logging
import re
import shutil
import subprocess
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

from backend import config

logger = logging.getLogger(__name__)


class ExtractionError(Exception):
    """A file could not be read. The message is safe to show to a student."""


@dataclass
class Slide:
    number: int  # 1-based, matches what students see in the deck
    title: str
    text: str
    notes: str = ""
    has_visual: bool = False

    @property
    def is_administrative(self) -> bool:
        return bool(
            re.match(r"^(thank\s+you|acknowledg[e]?ments?|references|credits)\b", self.title, re.I)
            or re.match(r"^(anchor:|slides adapted from)\s", self.text.strip(), re.I)
            or (
                self.number == 1
                and re.search(r"acknowledg[e]?ments?|department of|@\w+\.", self.text, re.I)
            )
        )

    @property
    def word_count(self) -> int:
        return len(f"{self.title} {self.text}".split())

    @property
    def is_image_heavy(self) -> bool:
        # Little text usually means a diagram or screenshot. These go to the
        # vision model in Stage 5; text-only summarization would miss them.
        return self.word_count < config.IMAGE_HEAVY_WORD_THRESHOLD


@dataclass
class Deck:
    title: str
    slides: list[Slide] = field(default_factory=list)
    vision_failures: list[int] = field(default_factory=list)


def extract_deck(path: str | Path) -> Deck:
    """Read any supported file type and return its slides."""
    path = Path(path)
    if not path.is_file():
        raise ExtractionError(f"File not found: {path}")

    suffix = path.suffix.lower()
    if suffix not in config.ALLOWED_EXTENSIONS:
        allowed = ", ".join(config.ALLOWED_EXTENSIONS)
        raise ExtractionError(f"Unsupported file type '{suffix}'. Allowed: {allowed}")

    if suffix == ".pptx":
        deck = _extract_pptx(path)
    elif suffix == ".ppt":
        deck = _extract_ppt(path)
    else:
        deck = _extract_pdf(path)

    if len(deck.slides) > config.MAX_SLIDES:
        raise ExtractionError(
            f"Deck has {len(deck.slides)} slides; the limit is {config.MAX_SLIDES}."
        )
    logger.info("Extracted %d slides from %s", len(deck.slides), path.name)
    return deck


# --- .pptx ---------------------------------------------------------------


def _extract_pptx(path: Path) -> Deck:
    try:
        prs = Presentation(str(path))
    except (PackageNotFoundError, KeyError, ValueError) as exc:
        raise ExtractionError(f"Could not open '{path.name}': it may be corrupt.") from exc

    slides = []
    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None else ""
        # Skip the title shape in the body so the title isn't repeated.
        # Compare shape_id, not identity: python-pptx builds a new proxy
        # object on every access, so `is` never matches.
        title_id = title_shape.shape_id if title_shape is not None else None
        body_shapes = [s for s in slide.shapes if s.shape_id != title_id]
        text = "\n".join(_shape_texts(body_shapes))

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        visual_types = {
            MSO_SHAPE_TYPE.PICTURE,
            MSO_SHAPE_TYPE.CHART,
            MSO_SHAPE_TYPE.GROUP,
            MSO_SHAPE_TYPE.FREEFORM,
        }
        has_visual = (
            any(
                s.shape_type in visual_types
                and s.width * s.height > prs.slide_width * prs.slide_height * 0.06
                for s in body_shapes
            )
            or sum(
                s.shape_type in {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.LINE}
                for s in body_shapes
            )
            >= 3
        )
        slides.append(
            Slide(number=number, title=title, text=text, notes=notes, has_visual=has_visual)
        )

    return Deck(title=_pick_title(prs.core_properties.title, slides, path.stem), slides=slides)


def _shape_texts(shapes) -> list[str]:
    """Collect text from text boxes, tables and grouped shapes, in slide order."""
    texts = []
    for shape in shapes:
        # Grouped shapes hide their text one level down, so recurse.
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(_shape_texts(shape.shapes))
        elif shape.has_text_frame:
            text = _frame_text(shape.text_frame)
            if text:
                texts.append(text)
        elif shape.has_table:
            # One line per row, cells separated by " | ", so the model can
            # still tell columns apart.
            for row in shape.table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    texts.append(" | ".join(cells))
    return texts


def _frame_text(text_frame) -> str:
    """A text frame's paragraphs, with superscript runs written as exponents
    (python-pptx's .text drops the raising, turning "2⁸¹" into "281")."""
    paragraphs = []
    for paragraph in text_frame.paragraphs:
        text = ""
        for run in paragraph.runs:
            # PowerPoint stores superscript as a positive baseline offset.
            baseline = run.font._rPr.get("baseline") if run.font._rPr is not None else None
            if baseline and int(baseline) > 0:
                text += _exponent(text, run.text)
            else:
                text += run.text
        paragraphs.append(text)
    return "\n".join(paragraphs).strip()


def _first_title(slides: list[Slide]) -> str:
    return next((s.title for s in slides if s.title), "")


# PowerPoint fills in placeholder metadata titles that are useless to students.
_GENERIC_TITLES = {"powerpoint presentation", "presentation", "slide 1", "untitled"}


def _pick_title(metadata_title: str | None, slides: list[Slide], fallback: str) -> str:
    title = (metadata_title or "").strip()
    if title.lower() in _GENERIC_TITLES:
        title = ""
    return title or _first_title(slides) or fallback


# --- .ppt (legacy binary format) -----------------------------------------


def _extract_ppt(path: Path) -> Deck:
    # python-pptx can't read the old binary .ppt format, so let LibreOffice
    # convert it to .pptx first and reuse the .pptx path.
    soffice = shutil.which("soffice")
    if soffice is None:
        raise ExtractionError(
            "Reading .ppt files needs LibreOffice. Install it with "
            "'brew install --cask libreoffice', or save the file as .pptx/.pdf."
        )

    with tempfile.TemporaryDirectory() as out_dir:
        try:
            subprocess.run(
                [soffice, "--headless", "--convert-to", "pptx", "--outdir", out_dir, str(path)],
                check=True,
                capture_output=True,
                timeout=120,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
            raise ExtractionError(f"Could not convert '{path.name}' from .ppt.") from exc

        converted = Path(out_dir) / f"{path.stem}.pptx"
        if not converted.exists():
            raise ExtractionError(f"Could not convert '{path.name}': it may be corrupt.")
        deck = _extract_pptx(converted)

    # The converted file lives in a temp dir, so fall back to the real name.
    if deck.title == converted.stem:
        deck.title = path.stem
    return deck


# --- .pdf ----------------------------------------------------------------


def _extract_pdf(path: Path) -> Deck:
    try:
        doc = pymupdf.open(path)
    except (pymupdf.FileDataError, RuntimeError) as exc:
        raise ExtractionError(f"Could not open '{path.name}': it may be corrupt.") from exc

    with doc:
        pages = [_page_lines(page) for page in doc]
        visuals = [_page_has_visual(page) for page in doc]
        meta_title = (doc.metadata or {}).get("title", "")

    boilerplate = _repeated_lines(pages)
    slides = []
    for number, lines in enumerate(pages, start=1):
        lines = [line for line in lines if line not in boilerplate]
        # PDFs have no "title" field per page; once headers/footers are gone,
        # the first line of a slide export is almost always the slide heading.
        title = lines[0] if lines else ""
        text = "\n".join(lines[1:])
        slides.append(Slide(number=number, title=title, text=text, has_visual=visuals[number - 1]))

    return Deck(title=_pick_title(meta_title, slides, path.stem), slides=slides)


def _page_has_visual(page):
    """A heuristic for substantive pictures/vector diagrams, not just logos.

    Keep the entire source page when selected so no arrow/label is cropped.
    Tables may also be selected, which is useful for visual reference.
    """
    area = page.rect.get_area()
    if any(pymupdf.Rect(image["bbox"]).get_area() > area * 0.08 for image in page.get_image_info()):
        return True
    body = pymupdf.Rect(0, page.rect.height * 0.15, page.rect.width, page.rect.height * 0.9)
    # Filled paths can be text glyphs/highlights. Count stroked connectors and
    # outlined shapes rather than mistaking ordinary text for diagrams.
    shapes = [
        d["rect"] for d in page.get_drawings() if "s" in d["type"] and body.contains(d["rect"])
    ]
    return len(shapes) >= 3 and sum(r.get_area() for r in shapes) > area * 0.04


def _page_lines(page: pymupdf.Page) -> list[str]:
    """Page text with tables kept as rows.

    Plain text extraction reads a table one cell per line ("sky", "temp", ...,
    "sunny", "warm", ...), so the model has to count positions to know which
    column a value is in, and it miscounts. Detected tables are instead
    written as "a | b | c" rows, like .pptx tables.
    """
    tables = _page_tables(page)

    # Text blocks outside the tables, in reading order.
    outside = []
    for block in page.get_text("dict")["blocks"]:
        x0, y0, x1, y1 = block["bbox"]
        centre = pymupdf.Point((x0 + x1) / 2, (y0 + y1) / 2)
        if block["type"] == 0 and not any(centre in rect for rect, _ in tables):
            outside.append(_block_text(block))
    lines = _clean_lines("\n".join(outside))

    for _, rows in tables:
        rows = [row for row in rows if any(row)]
        # Double borders between cells can be detected as extra, always-empty
        # columns; drop those.
        keep = [c for c in range(len(rows[0]) if rows else 0) if any(row[c] for row in rows)]
        lines += [" | ".join(row[c] for c in keep) for row in rows]
    return lines


Table = tuple[pymupdf.Rect, list[list[str]]]  # (area on the page, rows of cell text)


def _page_tables(page: pymupdf.Page) -> list[Table]:
    """Tables on the page, as cell text by row.

    Fast path: PowerPoint exports (like the college's PDFs) draw every table
    cell as its own rectangle, so a grid of touching rectangles is a table;
    that takes milliseconds. PyMuPDF's general find_tables() also handles
    tables drawn with plain lines, but costs ~26 s/page on a 0.1-CPU free
    host, so it's only the fallback for pages with a grid of drawn lines.
    """
    tables = _rectangle_tables(page)
    if tables or not config.SLOW_TABLE_FALLBACK or not _has_line_grid(page):
        return tables
    found = []
    for table in page.find_tables().tables:
        # Read each cell by its position: the table's own text extraction
        # scrambles letter-spaced text ("s y" for "sky") in some exports.
        rows = [
            [
                " ".join(page.get_textbox(pymupdf.Rect(cell)).split()) if cell else ""
                for cell in row.cells
            ]
            for row in table.rows
        ]
        found.append((pymupdf.Rect(table.bbox), rows))
    return found


_EDGE_TOLERANCE = 2.5  # points: edges this close count as touching or aligned


def _rectangle_tables(page: pymupdf.Page) -> list[Table]:
    words = page.get_text("words")  # (x0, y0, x1, y1, text, block, line, word)
    tables = []
    for cells in _touching_groups(_cell_rectangles(page)):
        rows = _distinct(cell.y0 for cell in cells)
        cols = _distinct(cell.x0 for cell in cells)
        # A table is a (nearly) full grid; scattered boxes in a layout aren't.
        if len(cells) < 4 or len(rows) < 2 or len(cols) < 2:
            continue
        if len(cells) < 0.75 * len(rows) * len(cols):
            continue
        grid = [["" for _ in cols] for _ in rows]
        for cell in cells:
            r = min(range(len(rows)), key=lambda i: abs(rows[i] - cell.y0))
            c = min(range(len(cols)), key=lambda i: abs(cols[i] - cell.x0))
            grid[r][c] = _words_in(cell, words)
        area = pymupdf.Rect(
            min(cell.x0 for cell in cells),
            min(cell.y0 for cell in cells),
            max(cell.x1 for cell in cells),
            max(cell.y1 for cell in cells),
        )
        tables.append((area, grid))
    return tables


def _cell_rectangles(page: pymupdf.Page) -> list[pymupdf.Rect]:
    """Drawn rectangles that could be table cells, minus overlays."""
    page_area = page.rect.width * page.rect.height
    candidates = [
        pymupdf.Rect(item[1])
        for drawing in page.get_drawings()
        for item in drawing["items"]
        if item[0] == "re"
    ]
    candidates = [
        r
        for r in candidates
        if r.width > 12 and r.height > 8 and r.width * r.height < page_area / 2
    ]
    # A rectangle covering smaller ones (a row highlight, or the same cell
    # drawn twice) isn't a cell. Going smallest first keeps the real cells
    # and skips the bigger box laid over them.
    cells: list[pymupdf.Rect] = []
    for rect in sorted(candidates, key=lambda r: r.width * r.height):
        if not any(_area(rect & cell) > 0.5 * _area(cell) for cell in cells):
            cells.append(rect)
    return cells


def _touching_groups(cells: list[pymupdf.Rect]) -> list[list[pymupdf.Rect]]:
    """Group rectangles that share an edge with a neighbour (flood fill)."""
    groups, seen = [], set()
    for start in range(len(cells)):
        if start in seen:
            continue
        seen.add(start)
        stack, group = [start], []
        while stack:
            i = stack.pop()
            group.append(cells[i])
            for j in range(len(cells)):
                if j not in seen and _share_edge(cells[i], cells[j]):
                    seen.add(j)
                    stack.append(j)
        groups.append(group)
    return groups


def _share_edge(a: pymupdf.Rect, b: pymupdf.Rect) -> bool:
    tol = _EDGE_TOLERANCE
    x_overlap = min(a.x1, b.x1) - max(a.x0, b.x0)
    y_overlap = min(a.y1, b.y1) - max(a.y0, b.y0)
    side_by_side = (abs(a.x1 - b.x0) < tol or abs(b.x1 - a.x0) < tol) and y_overlap > tol
    stacked = (abs(a.y1 - b.y0) < tol or abs(b.y1 - a.y0) < tol) and x_overlap > tol
    return side_by_side or stacked


def _distinct(values) -> list[float]:
    """Sorted coordinates, merging ones within the edge tolerance."""
    out: list[float] = []
    for value in sorted(values):
        if not out or value - out[-1] > _EDGE_TOLERANCE:
            out.append(value)
    return out


def _words_in(cell: pymupdf.Rect, words: list) -> str:
    """The words whose centre lies in the cell, in reading order."""
    inside = [
        w
        for w in words
        if cell.x0 - 1 <= (w[0] + w[2]) / 2 <= cell.x1 + 1
        and cell.y0 - 1 <= (w[1] + w[3]) / 2 <= cell.y1 + 1
    ]
    inside.sort(key=lambda w: (w[5], w[6], w[7]))  # block, line, word number
    return " ".join(w[4] for w in inside)


def _area(rect: pymupdf.Rect) -> float:
    return 0.0 if rect.is_empty else rect.width * rect.height


def _has_line_grid(page: pymupdf.Page) -> bool:
    """Cheap check (~1 ms/page): at least 3 horizontal and 3 vertical drawn
    lines, as a table ruled with lines (not cell rectangles) needs."""
    horizontal = vertical = 0
    for drawing in page.get_drawings():
        for item in drawing["items"]:
            if item[0] == "l":
                start, end = item[1], item[2]
                if abs(start.y - end.y) < 1 and abs(start.x - end.x) > 20:
                    horizontal += 1
                elif abs(start.x - end.x) < 1 and abs(start.y - end.y) > 10:
                    vertical += 1
            elif item[0] == "re":
                rect = item[1]
                if rect.width > 20 and rect.height < 2:  # a thin rectangle used as a line
                    horizontal += 1
                elif rect.height > 10 and rect.width < 2:
                    vertical += 1
    return horizontal >= 3 and vertical >= 3


_SUPERSCRIPT_FLAG = 1  # PyMuPDF span flag bit for superscript text


def _exponent(before: str, piece: str) -> str:
    """Write superscript text: "^81" after a digit, letter or ")", since
    that's an exponent; otherwise unchanged (the "ve" in "+ve")."""
    exponent = piece.strip()
    if not exponent or not (before[-1:].isalnum() or before[-1:] == ")"):
        return piece
    marked = f"^{exponent}" if exponent.isalnum() else f"^({exponent})"
    return marked + (" " if piece.endswith(" ") else "")


def _block_text(block: dict) -> str:
    """A text block's lines, with exponents kept as "^".

    Plain extraction drops the raising of superscripts, so "2⁸¹" on a slide
    became "281" and the model wrote "2^81 = 281". A superscript right after
    a digit, letter or ")" is an exponent and becomes "^81" (or "^(x+1)" if it
    isn't a single token). Other superscripts, like the "ve" in "+ve", are
    left as plain text.
    """
    lines = []
    for line in block["lines"]:
        text = ""
        for span in line["spans"]:
            piece = span["text"]
            if span["flags"] & _SUPERSCRIPT_FLAG:
                text += _exponent(text, piece)
            else:
                text += piece
        lines.append(text)
    return "\n".join(lines)


# Bullet glyphs PowerPoint exports as text, sometimes on a line of their own.
_BULLETS = "•·●■◆▪○◦➢➤✓"  # not "-": it could be a minus sign


def _clean_lines(page_text: str) -> list[str]:
    lines = []
    for raw in page_text.splitlines():
        line = raw.strip()
        if not line.strip(_BULLETS + " "):
            continue  # empty, or a bullet glyph with no text
        if line[0] in _BULLETS:
            line = "- " + line.lstrip(_BULLETS + " ")
        lines.append(line)
    return lines


def _repeated_lines(pages: list[list[str]]) -> set[str]:
    """Lines that appear on at least half the pages: running headers, footers,
    course names. They carry no slide content and confuse title detection."""
    if len(pages) < 4:  # too few pages to tell boilerplate from content
        return set()
    counts = Counter(line for lines in pages for line in set(lines))
    return {line for line, n in counts.items() if n >= len(pages) / 2}
