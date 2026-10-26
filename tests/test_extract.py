import shutil

import pytest

from backend.extract import ExtractionError, extract_deck


def test_pptx_title_text_and_notes(sample_pptx):
    deck = extract_deck(sample_pptx)

    assert deck.title == "Intro to Machine Learning"
    assert len(deck.slides) == 3

    first = deck.slides[0]
    assert first.number == 1
    assert first.title == "What is ML?"
    assert "labelled examples" in first.text
    assert "What is ML?" not in first.text  # title is not duplicated in body
    assert first.notes == "Mention spam filters as an example."


def test_pptx_tables_keep_columns(sample_pptx):
    table_slide = extract_deck(sample_pptx).slides[1]
    assert "Supervised | Regression" in table_slide.text


def test_image_heavy_detection(sample_pptx):
    slides = extract_deck(sample_pptx).slides
    assert not slides[0].is_image_heavy
    assert slides[2].is_image_heavy  # title only, no body text


def test_pdf_one_page_per_slide(sample_pdf):
    deck = extract_deck(sample_pdf)

    assert [s.title for s in deck.slides] == ["Linear Regression", "Gradient Descent"]
    assert "squared error" in deck.slides[0].text
    assert deck.title == "Linear Regression"  # no PDF metadata title, so first heading


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_legacy_ppt_via_libreoffice(sample_pptx, tmp_path):
    # Make a real .ppt by asking LibreOffice to convert our .pptx down.
    import subprocess

    subprocess.run(
        [
            "soffice",
            "--headless",
            "--convert-to",
            "ppt",
            "--outdir",
            str(tmp_path),
            str(sample_pptx),
        ],
        check=True,
        capture_output=True,
        timeout=120,
    )
    deck = extract_deck(tmp_path / "intro.ppt")
    assert len(deck.slides) == 3
    assert "labelled examples" in deck.slides[0].text


def test_rejects_unsupported_type(tmp_path):
    path = tmp_path / "notes.docx"
    path.write_text("hello")
    with pytest.raises(ExtractionError, match="Unsupported file type"):
        extract_deck(path)


@pytest.mark.parametrize("name", ["broken.pptx", "broken.pdf"])
def test_corrupt_file_gives_clear_error(tmp_path, name):
    path = tmp_path / name
    path.write_bytes(b"this is not a real file")
    with pytest.raises(ExtractionError, match="may be corrupt"):
        extract_deck(path)


def test_missing_file(tmp_path):
    with pytest.raises(ExtractionError, match="File not found"):
        extract_deck(tmp_path / "nope.pptx")


def test_pdf_strips_repeated_headers_and_bullet_glyphs(tmp_path):
    # Mimics the college exports: a course-name header on every page, above
    # the real slide heading, plus bullet glyphs exported as text.
    import pymupdf

    doc = pymupdf.open()
    for heading in ["Concept", "Feature Space", "Version Space", "Inductive Bias"]:
        page = doc.new_page()
        page.insert_text((72, 40), "Machine Learning", fontsize=10)
        page.insert_text((72, 80), heading, fontsize=20)
        # Body lines differ per page; identical ones would count as boilerplate.
        page.insert_text((72, 120), f"• point about {heading}", fontsize=12)
        page.insert_text((72, 140), f"-0.5 is a weight for {heading}", fontsize=12)
    path = tmp_path / "headers.pdf"
    doc.save(path)

    slides = extract_deck(path).slides
    assert [s.title for s in slides] == [
        "Concept",
        "Feature Space",
        "Version Space",
        "Inductive Bias",
    ]
    assert "Machine Learning" not in slides[0].text
    assert "- point about Concept" in slides[0].text
    assert "-0.5 is a weight for Concept" in slides[0].text  # minus sign untouched


def test_pdf_tables_become_rows(tmp_path):
    # A ruled table like the EnjoySport training data on the college slides.
    import pymupdf

    rows = [["sky", "water", "enjoy"], ["sunny", "warm", "yes"], ["sunny", "cool", "yes"]]
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 60), "Find S Algorithm - Example", fontsize=18)
    x0, y0, w, h = 72, 100, 120, 30
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            rect = pymupdf.Rect(x0 + c * w, y0 + r * h, x0 + (c + 1) * w, y0 + (r + 1) * h)
            page.draw_rect(rect, color=(0, 0, 0), width=1)
            page.insert_text((rect.x0 + 5, rect.y0 + 20), value, fontsize=11)
    path = tmp_path / "table.pdf"
    doc.save(path)

    slide = extract_deck(path).slides[0]
    assert slide.title == "Find S Algorithm - Example"
    assert "sky | water | enjoy" in slide.text
    assert "sunny | cool | yes" in slide.text
    assert "\nsky\n" not in f"\n{slide.text}\n"  # no one-cell-per-line leftovers


def test_pptx_superscripts_become_exponents(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Concept Space"
    paragraph = slide.placeholders[1].text_frame.paragraphs[0]
    for text, superscript in [
        ("#Concepts = 2", False),
        ("81", True),
        (" and +", False),
        ("ve", True),
    ]:
        run = paragraph.add_run()
        run.text = text
        if superscript:
            run.font._rPr.set("baseline", "30000")
    path = tmp_path / "sup.pptx"
    prs.save(path)

    text = extract_deck(path).slides[0].text
    assert "#Concepts = 2^81" in text
    assert "+ve" in text  # superscript after "+" isn't an exponent


def test_exponent_marking_rules():
    from backend.extract import _exponent

    assert _exponent("2", "81") == "^81"
    assert _exponent("2^(2", "d") == "^d"
    assert _exponent("(x+1)", "2 ") == "^2 "
    assert _exponent("e", "x+1") == "^(x+1)"
    assert _exponent("+", "ve") == "ve"
    assert _exponent("2", " ") == " "


def test_pages_without_ruling_lines_skip_table_detection(tmp_path, monkeypatch):
    import pymupdf

    calls = []
    real = pymupdf.Page.find_tables
    monkeypatch.setattr(
        pymupdf.Page, "find_tables", lambda self, *a, **k: calls.append(1) or real(self, *a, **k)
    )
    doc = pymupdf.open()
    for i in range(4):  # plain text slides, no drawn lines
        doc.new_page().insert_text((72, 72), f"Slide {i} heading", fontsize=18)
    path = tmp_path / "plain.pdf"
    doc.save(path)

    extract_deck(path)
    assert calls == []  # the expensive detector never ran


def _cell_grid_pdf(tmp_path, overlay=False):
    """A 3x3 table drawn PowerPoint-style: one rectangle per cell."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 60), "Training data", fontsize=18)
    values = [["sky", "water", "enjoy"], ["sunny", "warm", "yes"], ["sunny", "cool", "yes"]]
    for r, row in enumerate(values):
        for c, value in enumerate(row):
            cell = pymupdf.Rect(72 + 120 * c, 100 + 30 * r, 192 + 120 * c, 130 + 30 * r)
            page.draw_rect(cell, color=(0, 0, 0), fill=(0.9, 0.9, 1))
            page.insert_text((cell.x0 + 5, cell.y0 + 20), value, fontsize=11)
    if overlay:  # a highlight box laid over the whole second row
        page.draw_rect(pymupdf.Rect(70, 128, 434, 162), color=(1, 0, 0), width=2)
    path = tmp_path / "cells.pdf"
    doc.save(path)
    return path


def test_cell_rectangle_tables_are_read_without_the_slow_detector(tmp_path, monkeypatch):
    import pymupdf

    monkeypatch.setattr(pymupdf.Page, "find_tables", lambda *a, **k: pytest.fail("slow path used"))
    text = extract_deck(_cell_grid_pdf(tmp_path)).slides[0].text
    assert text.splitlines()[-3:] == [
        "sky | water | enjoy",
        "sunny | warm | yes",
        "sunny | cool | yes",
    ]


def test_highlight_box_over_a_row_is_not_a_cell(tmp_path):
    text = extract_deck(_cell_grid_pdf(tmp_path, overlay=True)).slides[0].text
    assert "sunny | warm | yes" in text


def test_scattered_boxes_are_not_a_table(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    # Two separate rows of boxes far apart: a layout, not a grid.
    for i, (x, y) in enumerate([(72, 100), (192, 100), (72, 400), (300, 600), (420, 600)]):
        page.draw_rect(pymupdf.Rect(x, y, x + 120, y + 30), color=(0, 0, 0))
        page.insert_text((x + 5, y + 20), f"box {i}", fontsize=11)
    path = tmp_path / "layout.pdf"
    doc.save(path)
    assert " | " not in extract_deck(path).slides[0].text


def test_slow_fallback_can_be_turned_off(tmp_path, monkeypatch):
    import pymupdf

    from backend import config

    doc = pymupdf.open()
    page = doc.new_page()
    for i in range(4):  # a grid ruled with plain lines
        page.draw_line((72, 100 + 30 * i), (432, 100 + 30 * i))
        page.draw_line((72 + 120 * i, 100), (72 + 120 * i, 190))
    path = tmp_path / "ruled.pdf"
    doc.save(path)

    monkeypatch.setattr(config, "SLOW_TABLE_FALLBACK", False)
    monkeypatch.setattr(pymupdf.Page, "find_tables", lambda *a, **k: pytest.fail("slow path used"))
    extract_deck(path)  # no failure: the slow detector never ran
