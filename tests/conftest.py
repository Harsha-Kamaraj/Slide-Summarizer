"""Fixtures that build tiny decks on the fly, so tests never depend on real
lecture files (which can't be committed)."""

from pathlib import Path

import pymupdf
import pytest
from pptx import Presentation
from pptx.util import Inches


@pytest.fixture
def sample_pptx(tmp_path: Path) -> Path:
    prs = Presentation()
    prs.core_properties.title = "Intro to Machine Learning"

    # Slide 1: title + bullets + speaker notes
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "What is ML?"
    slide.placeholders[1].text = (
        "Machine learning lets computers learn patterns from data\n"
        "Supervised learning uses labelled examples to train a model"
    )
    slide.notes_slide.notes_text_frame.text = "Mention spam filters as an example."

    # Slide 2: a table
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Types of learning"
    table = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, (a, b) in enumerate(
        [("Type", "Example"), ("Supervised", "Regression"), ("Unsupervised", "Clustering")]
    ):
        table.cell(r, 0).text = a
        table.cell(r, 1).text = b

    # Slide 3: almost no text (stands in for a diagram slide)
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Pipeline diagram"

    path = tmp_path / "intro.pptx"
    prs.save(path)
    return path


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    doc = pymupdf.open()
    for heading, body in [
        ("Linear Regression", "Fits a straight line to minimise squared error"),
        ("Gradient Descent", "Iteratively updates weights opposite to the gradient"),
    ]:
        page = doc.new_page()
        page.insert_text((72, 72), heading, fontsize=20)
        page.insert_text((72, 120), body, fontsize=12)
    path = tmp_path / "regression.pdf"
    doc.save(path)
    doc.close()
    return path
